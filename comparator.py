import io
from dataclasses import dataclass
from typing import Dict, List, Optional, Set, Tuple
import numpy as np
import openpyxl
from parser import HplcReport, Peak


@dataclass
class MasterImpurityRow:
    sr_no: int
    name: str
    mean_rt: float
    rrt: float
    is_main: bool
    batch_values: Dict[str, any]


@dataclass
class VerticalComparisonResult:
    rows: List[MasterImpurityRow]
    summary_rows: List[Dict[str, any]]
    batch_names: List[str]
    active_wavelength: Optional[int]
    available_wavelengths: List[int]
    main_peak_rts: Dict[str, float]


class HplcComparator:

    @staticmethod
    def parse_existing_vertical_excel(excel_bytes: bytes) -> Tuple[List[Dict[str, any]], List[str]]:
        """Parses an existing vertical comparison matrix (.xlsx) back into memory."""
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
        ws = wb.active

        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        existing_batch_names = [str(h).strip() for h in headers[4:] if h is not None and str(h).strip() != ""]

        existing_rows: List[Dict[str, any]] = []
        for r in range(2, ws.max_row + 1):
            name_val = ws.cell(row=r, column=2).value
            if name_val is None:
                continue

            name_str = str(name_val).strip()
            if any(k in name_str.lower() for k in ["total", "assay"]):
                continue

            rrt_val = ws.cell(row=r, column=4).value
            if rrt_val is None:
                continue

            try:
                rrt = round(float(rrt_val), 3)
                rt_val = ws.cell(row=r, column=3).value
                rt = round(float(rt_val), 3) if rt_val not in [None, "", "-"] else 0.0
            except ValueError:
                continue

            is_main = abs(rrt - 1.0) <= 0.008 or "api" in name_str.lower()

            b_vals: Dict[str, any] = {}
            for idx, b_name in enumerate(existing_batch_names, start=5):
                cell_val = ws.cell(row=r, column=idx).value
                if cell_val not in [None, ""]:
                    try:
                        b_vals[b_name] = float(cell_val)
                    except ValueError:
                        b_vals[b_name] = cell_val
                else:
                    b_vals[b_name] = ""

            existing_rows.append({
                "name": name_str,
                "mean_rt": rt,
                "rrt": rrt,
                "is_main": is_main,
                "batch_values": b_vals,
                "rts": [rt] if rt > 0 else [],
                "rrts": [rrt]
            })

        return existing_rows, existing_batch_names

    @classmethod
    def build_or_merge_vertical_matrix(
        cls,
        new_reports: List[HplcReport],
        existing_excel_bytes: Optional[bytes] = None,
        rrt_tolerance: float = 0.015,
        target_main_rt: Optional[float] = None,
        target_wavelength: Optional[int] = None,
        custom_specs: Optional[List[Dict[str, any]]] = None
    ) -> VerticalComparisonResult:
        existing_rows: List[Dict[str, any]] = []
        existing_batches: List[str] = []

        if existing_excel_bytes:
            try:
                existing_rows, existing_batches = cls.parse_existing_vertical_excel(existing_excel_bytes)
            except Exception:
                pass

        all_wl: Set[int] = set()
        for r in new_reports:
            all_wl.update(r.detected_wavelengths)
        available_wl_list = sorted(list(all_wl))

        active_wavelength = (
            target_wavelength
            if target_wavelength and target_wavelength > 0
            else (available_wl_list[0] if available_wl_list else None)
        )

        new_batch_names: List[str] = []
        for r in new_reports:
            b_id = r.batch_id
            if b_id in existing_batches or b_id in new_batch_names:
                count = 2
                while f"{b_id}_{count}" in existing_batches or f"{b_id}_{count}" in new_batch_names:
                    count += 1
                b_id = f"{b_id}_{count}"
            new_batch_names.append(b_id)

        all_batch_names = existing_batches + new_batch_names

        report_main_rts: Dict[str, float] = {}
        report_peaks_rrt: Dict[str, List[Tuple[Peak, float]]] = {}

        for b_name, r in zip(new_batch_names, new_reports):
            p_list = [
                p for p in r.peaks
                if active_wavelength is None or p.wavelength == 0 or p.wavelength == active_wavelength
            ]
            if not p_list:
                report_main_rts[b_name] = 0.0
                report_peaks_rrt[b_name] = []
                continue

            if target_main_rt and target_main_rt > 0:
                main_rt = min(p_list, key=lambda p: abs(p.retention_time - target_main_rt)).retention_time
            else:
                main_rt = max(p_list, key=lambda p: p.percent_area).retention_time
            report_main_rts[b_name] = main_rt

            rrt_items = []
            for p in p_list:
                if p.rel_rt is not None and p.rel_rt > 0:
                    c_rrt = round(p.rel_rt, 3)
                elif main_rt > 0:
                    c_rrt = round(p.retention_time / main_rt, 3)
                else:
                    c_rrt = 1.0
                rrt_items.append((p, c_rrt))
            report_peaks_rrt[b_name] = rrt_items

        master_rows: List[Dict[str, any]] = list(existing_rows)

        # Merge new batches into rows with RRT-priority greedy matching
        for b_name in new_batch_names:
            p_items = report_peaks_rrt.get(b_name, [])

            candidate_pairs = []
            for row in master_rows:
                for p, p_rrt in p_items:
                    diff = abs(row["rrt"] - p_rrt)
                    if diff <= rrt_tolerance:
                        candidate_pairs.append((diff, p, row))

            candidate_pairs.sort(key=lambda x: x[0])

            assigned_peaks = set()
            assigned_rows = set()

            for diff, p, row in candidate_pairs:
                peak_id = id(p)
                row_id = id(row)
                if peak_id not in assigned_peaks and row_id not in assigned_rows:
                    assigned_peaks.add(peak_id)
                    assigned_rows.add(row_id)
                    row["batch_values"][b_name] = 0.0 if p.percent_area == 0.0 else p.percent_area
                    row["rts"].append(p.retention_time)
                    row["rrts"].append(p.rel_rt if p.rel_rt else round(p.retention_time / report_main_rts[b_name], 3))
                    if (not row["name"] or row["name"] == "Unk") and p.name != "Unk":
                        row["name"] = p.name

            for p, p_rrt in p_items:
                if id(p) not in assigned_peaks:
                    is_main = abs(p_rrt - 1.0) <= 0.008
                    new_b_values = {eb: "" for eb in all_batch_names}
                    new_b_values[b_name] = 0.0 if p.percent_area == 0.0 else p.percent_area
                    master_rows.append({
                        "name": p.name if p.name != "Unk" else ("RIM (Main API)" if is_main else "Unk"),
                        "mean_rt": round(p.retention_time, 3),
                        "rrt": p_rrt,
                        "is_main": is_main,
                        "batch_values": new_b_values,
                        "rts": [p.retention_time],
                        "rrts": [p_rrt]
                    })

        for row in master_rows:
            for b in all_batch_names:
                if b not in row["batch_values"]:
                    row["batch_values"][b] = ""

        master_rows.sort(key=lambda x: x["rrt"])

        final_rows: List[MasterImpurityRow] = []
        for idx, item in enumerate(master_rows, start=1):
            m_rt = float(np.mean(item["rts"])) if item["rts"] else item["mean_rt"]
            m_rrt = float(np.mean(item["rrts"])) if item["rrts"] else item["rrt"]

            if item["is_main"]:
                m_rrt = 1.000

            resolved_name = item["name"]
            if custom_specs:
                for spec in custom_specs:
                    if spec["rrt_min"] <= m_rrt <= spec["rrt_max"]:
                        resolved_name = spec["name"]
                        break

            if item["is_main"] and (not resolved_name or resolved_name == "Unk"):
                resolved_name = "RIM (Main API)"

            final_rows.append(MasterImpurityRow(
                sr_no=idx,
                name=resolved_name,
                mean_rt=round(m_rt, 3),
                rrt=round(m_rrt, 3),
                is_main=item["is_main"],
                batch_values=item["batch_values"]
            ))

        total_imp = {"Name of Impurity": "Total Impurities (%)", "Mean RT (min)": "-", "RRT": "-"}
        main_assay = {"Name of Impurity": "Main Peak / Assay (%)", "Mean RT (min)": "-", "RRT": "1.000"}
        total_area = {"Name of Impurity": "Total Area (%)", "Mean RT (min)": "-", "RRT": "-"}

        for b_name in all_batch_names:
            imp_sum = 0.0
            api_val = 0.0
            for r in final_rows:
                v = r.batch_values.get(b_name, "")
                if v != "" and v is not None:
                    try:
                        num = float(v)
                        if r.is_main:
                            api_val += num
                        else:
                            imp_sum += num
                    except ValueError:
                        pass
            total_imp[b_name] = round(imp_sum, 2)
            main_assay[b_name] = round(api_val, 2)
            total_area[b_name] = round(imp_sum + api_val, 2)

        summary_rows = [total_imp, main_assay, total_area]

        return VerticalComparisonResult(
            rows=final_rows,
            summary_rows=summary_rows,
            batch_names=all_batch_names,
            active_wavelength=active_wavelength,
            available_wavelengths=available_wl_list,
            main_peak_rts=report_main_rts
        )
